# routers/manage.py — 数据管理域：班级/管理员账号、成绩 xlsx 导入、第二课堂导入、批量调分

import json
import re
import shutil
import uuid
from datetime import datetime

from fastapi import APIRouter, File, Form, Header, HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel

import db
from common import (
    ADJUST_FIELDS,
    ADJUST_OPS,
    DATA_DIR,
    MAX_FILE_SIZE,
    TOKENS,
    UPLOADS_DIR,
    check_class_scope,
    require_root,
    require_token,
)
from scoring import XLSX_PARSE_ERRORS, calc_total, drop_tmp, parse_xlsx

router = APIRouter()


class ClassNameBody(BaseModel):
    name: str


class AdminBody(BaseModel):
    username: str
    password: str
    class_id: str = ""
    role: str = "admin"


class AdminPasswordBody(BaseModel):
    password: str


class AdjustBody(BaseModel):
    sids: str
    field: str
    op: str
    points: float


@router.get("/api/classes")
def list_classes():
    return {"classes": db.list_classes()}


@router.post("/api/classes")
def add_class(body: ClassNameBody, authorization: str = Header(default="")):
    require_root(authorization)
    name = body.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="班级名称不能为空")
    if len(name) > 64:
        raise HTTPException(status_code=400, detail="班级名称过长")
    if any(c["name"] == name for c in db.list_classes()):
        raise HTTPException(status_code=400, detail="班级名称已存在")
    return {"ok": True, "class": db.create_class(name)}


@router.delete("/api/classes/{cid}")
def remove_class(cid: str, authorization: str = Header(default="")):
    require_root(authorization)
    if not db.get_class(cid):
        raise HTTPException(status_code=404, detail="班级不存在")
    if db.list_students(cid):
        raise HTTPException(status_code=400, detail="该班级仍有学生，请先清除班级数据")
    db.delete_class(cid)
    return {"ok": True}


@router.post("/api/classes/{cid}/clear")
def clear_class(cid: str, authorization: str = Header(default="")):
    require_root(authorization)
    cls = db.get_class(cid)
    if not cls:
        raise HTTPException(status_code=404, detail="班级不存在")
    folders = db.clear_class_data(cid)
    db.set_class_meta(cid, 0, "")
    for folder in folders:
        if folder and not db.folder_in_use(folder):
            shutil.rmtree(UPLOADS_DIR / folder, ignore_errors=True)
    return {"ok": True, "class": cls["name"]}


@router.post("/api/admins")
def add_admin(body: AdminBody, authorization: str = Header(default="")):
    require_root(authorization)
    username = body.username.strip()
    if not re.fullmatch(r"[A-Za-z0-9_]{3,32}", username):
        raise HTTPException(status_code=400, detail="用户名限 3-32 位字母/数字/下划线")
    if username == "root" or db.get_admin(username):
        raise HTTPException(status_code=400, detail="用户名已存在")
    if len(body.password) < 6:
        raise HTTPException(status_code=400, detail="密码至少 6 位")
    if body.role != "admin":
        raise HTTPException(status_code=400, detail="只能创建管理员账号")
    cls = db.get_class(body.class_id or "")
    if not cls:
        raise HTTPException(status_code=400, detail="负责的班级不存在")
    db.create_admin(username, body.password, "admin", cls["id"])
    return {"ok": True}


@router.delete("/api/admins/{username}")
def remove_admin(username: str, authorization: str = Header(default="")):
    session = require_root(authorization)
    admin = db.get_admin(username)
    if admin is None:
        raise HTTPException(status_code=404, detail="账号不存在")
    if admin["role"] == "root" or username == session["username"]:
        raise HTTPException(status_code=400, detail="不能删除该账号")
    db.delete_admin(username)
    for t, s in list(TOKENS.items()):
        if s["username"] == username:
            TOKENS.pop(t, None)
    return {"ok": True}


@router.put("/api/admins/{username}")
def reset_admin_password(username: str, body: AdminPasswordBody, authorization: str = Header(default="")):
    require_root(authorization)
    admin = db.get_admin(username)
    if admin is None:
        raise HTTPException(status_code=404, detail="账号不存在")
    if admin["role"] == "root":
        raise HTTPException(status_code=400, detail="超级管理员密码请由运维直接重置")
    if len(body.password) < 6:
        raise HTTPException(status_code=400, detail="密码至少 6 位")
    db.update_admin(username, password=body.password)
    return {"ok": True}


@router.get("/api/admins")
def get_admins(authorization: str = Header(default="")):
    require_root(authorization)
    return {"admins": db.list_admins()}


@router.post("/api/upload")
def upload(
    file: UploadFile = File(...),
    class_id: str = Form(""),
    authorization: str = Header(default=""),
):
    session = require_token(authorization)
    if session["role"] == "admin":
        class_id = session.get("class_id") or ""
        if not class_id:
            raise HTTPException(status_code=400, detail="你尚未分配负责班级，请联系 root")
    cls = db.get_class(class_id)
    if not cls:
        raise HTTPException(status_code=400, detail="班级不存在，请先创建班级")
    filename = file.filename or "upload.xlsx"
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 文件")
    tmp = DATA_DIR / f"{uuid.uuid4().hex}.xlsx"
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    with tmp.open("wb") as f:
        shutil.copyfileobj(file.file, f)
    try:
        students, rows = parse_xlsx(tmp)
    except HTTPException:
        tmp.unlink(missing_ok=True)
        raise
    except XLSX_PARSE_ERRORS:
        # 损坏/改名的非法 xlsx：给出可读中文提示，避免 500
        drop_tmp(tmp)
        raise HTTPException(status_code=400, detail="文件无法解析，请上传成绩导出 xlsx")
    db.replace_students(students, class_id)
    db.set_class_meta(class_id, rows, filename)
    tmp.unlink(missing_ok=True)
    return {"ok": True, "students": len(students), "rows": rows, "source": filename}


@router.post("/api/secondclass/import")
def import_secondclass(
    file: UploadFile = File(...),
    class_id: str = Form(""),
    threshold: float = Form(2.0),
    authorization: str = Header(default=""),
):
    """导入第二课堂统计 xlsx：对「学分 ≥ threshold」的学生德育(deyu)+10。

    - 防重复：同班再次导入先撤销上一次导入加的 10 分，再按新文件应用。
    - 未达标 / 名单外学生不动（保留现有德育分）。
    - 每次加减写 adjust_log 留痕；并为达标学生生成自动通过的德育申报记录
      （申报列表可见，meta 记录这批记录 id，下次导入按其撤销）。
    """
    session = require_token(authorization)
    if session["role"] == "admin":
        class_id = session.get("class_id") or ""
        if not class_id:
            raise HTTPException(status_code=400, detail="你尚未分配负责班级，请联系 root")
    elif not class_id:
        raise HTTPException(status_code=400, detail="请选择班级")
    check_class_scope(session, class_id)
    if not db.get_class(class_id):
        raise HTTPException(status_code=400, detail="班级不存在，请先创建班级")
    if threshold < 0:
        raise HTTPException(status_code=400, detail="阈值不能为负")
    filename = file.filename or "secondclass.xlsx"
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 文件")
    # meta.k 列仅 32 字符，class_id 32 位 hex，短 key 用前缀+截断
    meta_key = f"sc2_{class_id[:27]}"

    tmp = DATA_DIR / f"{uuid.uuid4().hex}.xlsx"
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    size = 0
    with tmp.open("wb") as f:
        while True:
            chunk = file.file.read(1024 * 1024)
            if not chunk:
                break
            size += len(chunk)
            if size > MAX_FILE_SIZE:
                f.close()
                tmp.unlink(missing_ok=True)
                raise HTTPException(status_code=400, detail="文件超过大小限制（10MB）")
            f.write(chunk)
    try:
        wb = load_workbook(tmp, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            raise HTTPException(status_code=400, detail="文件为空")
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        sid_col = next((i for i, h in enumerate(header) if h == "学号"), None)
        cred_col = next((i for i, h in enumerate(header) if h == "学分"), None)
        if sid_col is None or cred_col is None:
            raise HTTPException(
                status_code=400,
                detail=f"表头缺少「学号」或「学分」列（实际表头：{'、'.join(header[:8])}…）",
            )
        # 解析文件：sid -> 学分（单元格可能是文本，容错转换）
        file_credits = {}
        for r in rows[1:]:
            if not r or sid_col >= len(r) or cred_col >= len(r):
                continue
            sid = str(r[sid_col]).strip() if r[sid_col] is not None else ""
            if not sid:
                continue
            try:
                val = float(str(r[cred_col]).strip())
            except (TypeError, ValueError):
                continue
            file_credits[sid] = val
        if not file_credits:
            raise HTTPException(status_code=400, detail="未能从文件中解析出学号与学分数据")
    except XLSX_PARSE_ERRORS:
        # 损坏/改名的非法 xlsx：给出可读中文提示，避免 500
        raise HTTPException(status_code=400, detail="文件无法解析，请上传第二课堂统计 xlsx")
    finally:
        drop_tmp(tmp)

    pool = {s["sid"]: s for s in db.list_students(class_id)}
    now = datetime.now().isoformat(timespec="seconds")

    # 1) 防重复：撤销上一次导入。新格式 meta 带 aids（上次导入生成的申报记录），按记录撤销——
    #    已被管理员撤回/删除的记录不再重复扣分；旧格式仅记达标学号集，按 10 分撤销。
    ops = []  # {sid,name,op,points,old,new,total}
    delete_aids = []
    last_raw = db.get_meta(meta_key, "")
    last = {}
    if last_raw:
        try:
            last = json.loads(last_raw)
        except (json.JSONDecodeError, TypeError):
            last = {}
    if "aids" in last:
        for aid in last.get("aids") or []:
            rec = db.get_award(aid)
            if not rec:
                continue
            stu = pool.get(rec["sid"])
            if rec["approved"] == "是" and stu:
                old = round(float(stu["deyu"]), 2)
                new = round(max(old - float(rec["points"]), 0.0), 2)
                if new != old:
                    stu["deyu"] = new
                    total = calc_total(stu["deyu"], stu["score"], stu["tiyu"], stu["meiyu"], stu["laoyu"], stu["fujia"])
                    ops.append(
                        {
                            "sid": rec["sid"], "name": stu["name"], "op": "sub", "points": float(rec["points"]),
                            "old": old, "new": new, "total": total,
                        }
                    )
            delete_aids.append(aid)
    else:
        for sid in last.get("sids", []) or []:
            stu = pool.get(sid)
            if not stu:
                continue
            old = round(float(stu["deyu"]), 2)
            new = round(max(old - 10.0, 0.0), 2)
            if new == old:
                continue
            stu["deyu"] = new
            total = calc_total(stu["deyu"], stu["score"], stu["tiyu"], stu["meiyu"], stu["laoyu"], stu["fujia"])
            ops.append(
                {
                    "sid": sid, "name": stu["name"], "op": "sub", "points": 10.0,
                    "old": old, "new": new, "total": total,
                }
            )

    # 2) 应用新文件：学分 >= threshold 且在本班榜单的学生 +10（封顶 100），
    #    并为每个达标学生生成一条自动通过的德育申报记录（申报列表可见）
    applied = []
    skipped = []
    qualified = []
    records = []
    record_time = datetime.now()
    for sid, credit in file_credits.items():
        stu = pool.get(sid)
        if not stu:
            skipped.append({"sid": sid, "reason": "学号不在本班榜单"})
            continue
        if credit < threshold:
            continue  # 未达标不动
        qualified.append(sid)
        old = round(float(stu["deyu"]), 2)
        new = round(min(old + 10.0, 100.0), 2)
        if new != old:
            stu["deyu"] = new
            total = calc_total(stu["deyu"], stu["score"], stu["tiyu"], stu["meiyu"], stu["laoyu"], stu["fujia"])
            ops.append(
                {
                    "sid": sid, "name": stu["name"], "op": "add", "points": 10.0,
                    "old": old, "new": new, "total": total,
                }
            )
        records.append(
            {
                "id": uuid.uuid4().hex,
                "sid": sid,
                "name": stu["name"],
                "class_id": class_id,
                "category": "德育",
                "points": 10.0,
                "basis": f"第二课堂学分达标（{credit:g} 学分 ≥ {threshold:g}）",
                "evidence": [],
                "folder": "",
                "approved": "是",
                "created_at": record_time,
            }
        )
        applied.append({"sid": sid, "name": stu["name"], "credit": credit, "old": old, "new": new})

    # 3) 单事务写库（调分+留痕+删旧记录+写申报记录）+ 更新 meta
    meta_value = json.dumps(
        {
            "threshold": round(float(threshold), 2),
            "sids": qualified,
            "time": now,
            "aids": [r["id"] for r in records],
        },
        ensure_ascii=False,
    )
    if ops or records or delete_aids:
        for o in ops:
            o["created_at"] = now
        db.apply_secondclass(class_id, ops, meta_key, meta_value, records, delete_aids)
    else:
        db.set_meta(meta_key, meta_value)  # 无变化也更新阈值/名单，保证下次撤销依据最新

    return {
        "ok": True,
        "threshold": round(float(threshold), 2),
        "qualified": len(qualified),
        "applied": len(applied),
        "changes": len(ops),
        "skipped": skipped,
    }


@router.post("/api/scores/adjust")
def adjust_scores(body: AdjustBody, authorization: str = Header(default="")):
    session = require_token(authorization)
    if body.field not in ADJUST_FIELDS:
        raise HTTPException(status_code=400, detail="分数项无效，可选：" + "、".join(ADJUST_FIELDS.values()))
    if body.op not in ADJUST_OPS:
        raise HTTPException(status_code=400, detail="操作无效，可选 add（增加）/sub（减少）/set（设为）")
    if body.points < 0:
        raise HTTPException(status_code=400, detail="数值不能为负")
    cap = 5.0 if body.field == "fujia" else 100.0
    pool = db.list_students(session.get("class_id") if session["role"] == "admin" else None)
    tokens = [t for t in re.split(r"[\s,，;；]+", body.sids.strip()) if t]
    if not tokens:
        raise HTTPException(status_code=400, detail="请填写学号（支持正则，如 251184Y3.*）")
    matched = {}
    for token in tokens:
        try:
            rx = re.compile(token)
        except re.error as e:
            raise HTTPException(status_code=400, detail=f"无效的正则表达式 {token}：{e}")
        hits = [s for s in pool if s["sid"] == token or rx.fullmatch(s["sid"])]
        if not hits:
            raise HTTPException(status_code=400, detail=f"学号/正则 {token} 未匹配到任何学生")
        for s in hits:
            matched[s["sid"]] = s
    field = body.field
    op = body.op
    changes = []
    log_entries = []
    for stu in sorted(matched.values(), key=lambda s: s["sid"]):
        old = stu[field]
        if op == "add":
            new = min(old + body.points, cap)
        elif op == "sub":
            new = max(old - body.points, 0.0)
        else:
            new = min(max(body.points, 0.0), cap)
        new = round(new, 2)
        stu[field] = new
        stu["total"] = calc_total(stu["deyu"], stu["score"], stu["tiyu"], stu["meiyu"], stu["laoyu"], stu["fujia"])
        db.update_student_score(stu["sid"], stu["class_id"], field, new, stu["total"])
        changes.append({"sid": stu["sid"], "name": stu["name"], "field": field, "old": old, "new": new, "total": stu["total"]})
        log_entries.append(
            {
                "created_at": datetime.now().isoformat(timespec="seconds"),
                "sid": stu["sid"],
                "name": stu["name"],
                "field": field,
                "op": op,
                "points": body.points,
                "old": old,
                "new": new,
            }
        )
    db.insert_adjust_log(log_entries)
    return {"changed": changes}
