# routers/board.py — 榜单展示与导出：按班榜单、综合测评 Excel、证据 zip 归档

import io
from datetime import datetime
from urllib.parse import quote

from fastapi import APIRouter, Header, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import db
from common import check_class_scope, require_token
from evidence_zip import build_evidence_zip

router = APIRouter()


@router.get("/api/leaderboard")
def leaderboard(class_id: str | None = None):
    # 榜单必须指明班级：不带 class_id 会把多个班的学生按总分混排、统一编号名次，故直接拒绝
    cid = (class_id or "").strip()
    if not cid:
        raise HTTPException(status_code=400, detail="缺少 class_id：榜单按班级展示，请指定班级")
    students = db.list_students(cid)
    cls = db.get_class(cid)
    meta = {
        "rows": cls["row_count"] if cls else 0,
        "source": cls["source"] if cls else "",
        # 回传班级标识，便于前端/排查时自证当前展示的是哪个班（class_id 不存在时为空）
        "class_id": cid,
        "class_name": cls["name"] if cls else "",
    }
    # 挂科（无参评资格）学生沉底且不占名次：rank 置 0，综合测评成绩列显示 -1
    ranked = [s for s in students if not s.get("failed")]
    disq = [s for s in students if s.get("failed")]
    for rank, s in enumerate(ranked, start=1):
        s["rank"] = rank
    for s in disq:
        s["rank"] = 0
        s["total"] = -1
    return {"students": ranked + disq, "meta": meta}


@router.get("/api/export")
def export(class_id: str | None = None):
    # 榜单按班级导出：不带 class_id 会导出跨班混排的合并榜单（且文件名班级名为空），故直接拒绝
    cid = (class_id or "").strip()
    if not cid:
        raise HTTPException(status_code=400, detail="缺少 class_id：请先选择要导出的班级")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    headers = ["学号", "姓名", "德*15%", "智*60%", "体*10%", "美*5%", "劳*10%", "附加", "总分"]

    # 文件名对齐附件「25-26下学期xx班综合测评总分.xlsx」：xx 处放实际班级名
    cls = db.get_class(cid)
    cls_name = cls["name"] if cls else ""
    file_name = f"25-26下学期{cls_name}班综合测评总分.xlsx"

    # 样式严格对齐附件「25-26下学期xx班综合测评总分.xlsx」：
    # 等线 11 号（学号列仿宋）、全表细边框、学号列白色实底+文本格式、表头居中
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    font_body = Font(name="等线", size=11)
    font_sid = Font(name="仿宋", size=11)
    red_font = Font(name="等线", size=11, color="FFFF0000")
    red_sid_font = Font(name="仿宋", size=11, color="FFFF0000")
    center = Alignment(horizontal="center", vertical="center")
    header_font = Font(name="等线", size=11)

    ws.append(headers)
    for col in "ABCDEFGHI":
        cell = ws[f"{col}1"]
        cell.font = header_font
        cell.alignment = center
        cell.border = border
    # 学号列表头也用仿宋（附件 A1 是仿宋）
    ws["A1"].font = Font(name="仿宋", size=11)
    ws["A1"].fill = PatternFill(fill_type="solid", start_color="FFFFFFFF", end_color="FFFFFFFF")
    ws["A1"].number_format = "@"
    ws["B1"].number_format = "@"

    for stu in db.list_students(cid):
        total = -1 if stu.get("failed") else stu["total"]
        row = [
            str(stu["sid"]),
            stu["name"],
            stu["deyu"],
            stu["score"],
            stu["tiyu"],
            stu["meiyu"],
            stu["laoyu"],
            stu["fujia"],
            total,
        ]
        ws.append(row)
        r = ws.max_row
        for i, col in enumerate("ABCDEFGHI", start=1):
            cell = ws.cell(row=r, column=i)
            cell.border = border
            # 挂科（无参评资格）行整行标红
            cell.font = red_sid_font if (col == "A" and stu.get("failed")) else (red_font if stu.get("failed") else (font_sid if col == "A" else font_body))
            if col in ("A", "B"):
                # 学号/姓名：文本格式（附件 A、B 列均为 @），学号列附白色实底
                cell.number_format = "@"
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if col == "A":
                    cell.fill = PatternFill(fill_type="solid", start_color="FFFFFFFF", end_color="FFFFFFFF")
            elif i == 2:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(vertical="center")

    # 表头行高 + 列宽（附件值）
    ws.row_dimensions[1].height = 24.85
    for col, width in zip("ABC", (8.71, 8.71, 8.71)):
        ws.column_dimensions[col].width = width
    for col in "DEFGHI":
        ws.column_dimensions[col].width = 13
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename*=UTF-8''" + quote(file_name)},
    )


@router.get("/api/export/evidence-zip")
def export_evidence_zip(class_id: str | None = None, authorization: str = Header(default="")):
    session = require_token(authorization)
    if class_id:
        check_class_scope(session, class_id)
    else:
        if session["role"] != "root":
            class_id = session.get("class_id")
            if not class_id:
                raise HTTPException(status_code=403, detail="账号未绑定班级，无法导出")
    buf = build_evidence_zip(class_id)
    zip_name = f"evidence_backup_{datetime.now():%Y%m%d}.zip"
    return StreamingResponse(
        buf,
        media_type="application/zip",
        headers={
            "Content-Disposition": f"attachment; filename={zip_name}",
            # zip 已在内存完整生成，显式给 Content-Length，避免 chunked 下载在某些浏览器被拦截
            "Content-Length": str(len(buf.getvalue())),
        },
    )
