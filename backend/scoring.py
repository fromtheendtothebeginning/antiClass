# scoring.py — 智育/综合测评成绩计算与成绩 xlsx 解析（纯计算，无路由）

import zipfile
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

EXAM_PRIORITY = {"重修": 3, "补考一": 2, "正常考试": 1}


def is_excluded(row):
    category = row[21]
    course_name = row[6]
    invalid = row[25]
    if category == "通识课":
        return True
    if "体育" in course_name:
        return True
    if invalid == "是":
        return True
    return False


def to_float(value):
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return 0.0


def pick_best(course_rows):
    return max(course_rows, key=lambda r: (EXAM_PRIORITY.get(r[18], 1), to_float(r[3])))


def _first_exam_score(course_rows):
    """按办法「第一次考试」口径取课程成绩：优先正常考试（取其中最高），
    无正常考试记录才降级取 重修>补考一 的优先级最高记录；旷考按 0 分。"""
    normals = [r for r in course_rows if r[18] == "正常考试"]
    if normals:
        return max(to_float(r[3]) for r in normals)
    best = pick_best(course_rows)
    score = to_float(best[3])
    if best[17] == "旷考":
        score = 0.0
    return score


def compute_students(rows):
    groups = {}
    sports = {}
    names = {}
    for row in rows:
        sid = row[1]
        if not sid:
            continue
        names[sid] = row[4]
        if "体育" in row[6]:
            sports.setdefault((sid, row[2]), []).append(row)
        elif not is_excluded(row):
            groups.setdefault((sid, row[2]), []).append(row)

    students = {}
    for (sid, course_code), course_rows in groups.items():
        best = pick_best(course_rows)
        score = to_float(best[3])
        if best[17] == "旷考":
            score = 0.0
        credits = to_float(best[16])
        stu = students.setdefault(
            sid,
            {"sid": sid, "name": names[sid], "course_count": 0, "credits": 0.0, "score_weight": 0.0, "gpa_weight": 0.0, "tiyu_total": 0.0, "tiyu_count": 0, "failed": False},
        )
        stu["course_count"] += 1
        stu["credits"] += credits
        stu["score_weight"] += score * credits
        stu["gpa_weight"] += to_float(best[19]) * credits
        # 参评资格：当学期课程（不含通识课/体育课）第一次考试无不及格；任一 <60 即挂科
        if _first_exam_score(course_rows) < 60:
            stu["failed"] = True

    for (sid, course_code), sport_rows in sports.items():
        best = pick_best(sport_rows)
        stu = students.setdefault(
            sid,
            {"sid": sid, "name": names[sid], "course_count": 0, "credits": 0.0, "score_weight": 0.0, "gpa_weight": 0.0, "tiyu_total": 0.0, "tiyu_count": 0},
        )
        stu["tiyu_total"] += to_float(best[3])
        stu["tiyu_count"] += 1

    result = []
    for stu in students.values():
        credits = stu["credits"]
        zhiyu = round(stu["score_weight"] / credits, 2) if credits else 0.0
        gpa = round(stu["gpa_weight"] / credits, 2) if credits else 0.0
        tiyu = round(stu["tiyu_total"] / stu["tiyu_count"], 2) if stu["tiyu_count"] else 60.0
        deyu, meiyu, laoyu, fujia = 70.0, 70.0, 70.0, 0.0
        result.append(
            {
                "sid": stu["sid"],
                "name": stu["name"],
                "course_count": stu["course_count"],
                "credits": round(credits, 2),
                "gpa": gpa,
                "deyu": deyu,
                "score": zhiyu,
                "tiyu": tiyu,
                "meiyu": meiyu,
                "laoyu": laoyu,
                "fujia": fujia,
                "failed": stu["failed"],
                "total": calc_total(deyu, zhiyu, tiyu, meiyu, laoyu, fujia),
            }
        )
    result.sort(key=lambda s: (-s["total"], -s["score"], -s["gpa"], s["sid"]))
    for rank, stu in enumerate(result, start=1):
        stu["rank"] = rank
    return result


def calc_total(deyu, zhiyu, tiyu, meiyu, laoyu, fujia):
    return round(deyu * 0.15 + zhiyu * 0.60 + tiyu * 0.10 + meiyu * 0.05 + laoyu * 0.10 + fujia, 2)


# 解析 xlsx 时的“文件本身有问题”类异常：非 zip/损坏（BadZipFile）、
# 不支持的格式（InvalidFileException）、缺内部部件（KeyError）、空表/短行（IndexError）
XLSX_PARSE_ERRORS = (zipfile.BadZipFile, InvalidFileException, KeyError, IndexError)


def drop_tmp(path):
    """删除临时 xlsx。解析中途抛错时 openpyxl 可能仍占用文件句柄，Windows 会拒绝删除
    （PermissionError），此处忽略删除失败，避免把 400 顺手变成 500。"""
    try:
        path.unlink(missing_ok=True)
    except OSError:
        pass


def parse_xlsx(path):
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = rows[0]
    if "学号" not in header:
        raise HTTPException(status_code=400, detail="表头不含「学号」列，不是有效的成绩导出文件")
    data_rows = [row for row in rows[1:] if row and row[1]]
    return compute_students(data_rows), len(data_rows)
